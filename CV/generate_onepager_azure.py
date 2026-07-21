from pptx import Presentation
from pptx.util import Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation("/Users/karthiksankaran/PS-Scripts/CV/Karthik Sankaran-one-Pager.pptx")
slide = prs.slides[0]

# Slide dimensions: 25.40 cm wide x 19.05 cm tall
W = 25.40
H = 19.05
LC = 8.5   # left column width
RC_X = 8.8  # right column start x
RC_W = W - RC_X - 0.2  # right column width

NAVY   = RGBColor(0x1b, 0x4f, 0x72)
AZURE  = RGBColor(0x00, 0x78, 0xD4)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GREY   = RGBColor(0x55, 0x55, 0x55)
GREEN  = RGBColor(0x1e, 0x84, 0x49)
LTBLUE = RGBColor(0xAE, 0xD6, 0xF1)
BGBLUE = RGBColor(0xe8, 0xf4, 0xfd)

def set_tf(tf, text, size, bold=False, color=None, align=PP_ALIGN.LEFT, italic=False):
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = 'Calibri'
    if color:
        run.font.color.rgb = color

def add_textbox(slide, left, top, width, height, text, size, bold=False, color=None,
                align=PP_ALIGN.LEFT, italic=False):
    txBox = slide.shapes.add_textbox(Cm(left), Cm(top), Cm(width), Cm(height))
    txBox.word_wrap = True
    tf = txBox.text_frame
    tf.word_wrap = True
    set_tf(tf, text, size, bold, color, align, italic)
    return txBox

def add_rect(slide, left, top, width, height, fill_color):
    shape = slide.shapes.add_shape(1, Cm(left), Cm(top), Cm(width), Cm(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape

def add_bullet_box(slide, left, top, width, height, items, size=9, color=None):
    txBox = slide.shapes.add_textbox(Cm(left), Cm(top), Cm(width), Cm(height))
    txBox.word_wrap = True
    tf = txBox.text_frame
    tf.word_wrap = True
    c = color or RGBColor(0x1a, 0x1a, 0x1a)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(1)
        run = p.add_run()
        run.text = f'▸  {item}'
        run.font.size = Pt(size)
        run.font.name = 'Calibri'
        run.font.color.rgb = c

# ── Clear existing shapes ─────────────────────────────────────────────────────
for shape in list(slide.shapes):
    shape._element.getparent().remove(shape._element)

HEADER_H = 3.6

# ── HEADER BAR ────────────────────────────────────────────────────────────────
add_rect(slide, 0, 0, W, HEADER_H, AZURE)
add_rect(slide, 0, 0, W, 0.2, NAVY)

# Name
add_textbox(slide, 0.4, 0.3, 16, 1.2, 'Karthik Sankaran', 24, bold=True, color=WHITE)

# Title
add_textbox(slide, 0.4, 1.6, 17, 0.65,
            'Senior Manager  |  Associate Director  |  Azure Cloud & Infrastructure Lead',
            10, color=LTBLUE)

# Contact (right)
add_textbox(slide, 17.5, 0.35, 7.6, 1.0,
            'karthik.a.sankaran@outlook.com\n+44 747 130 2755  |  UK',
            8.5, color=LTBLUE, align=PP_ALIGN.RIGHT)

# Headline
add_textbox(slide, 0.4, 2.35, W - 0.8, 1.1,
            'Azure-specialised cloud leader · 30+ years experience · Enterprise migrations, landing zones, hybrid cloud & DevOps '
            'for global FS, Manufacturing & Media clients · £1.2M Azure savings · 15,000-user rollout · Accenture & Deloitte',
            8.5, color=WHITE)

# ── LEFT COLUMN ───────────────────────────────────────────────────────────────
add_rect(slide, 0, HEADER_H, LC, H - HEADER_H, BGBLUE)

y = HEADER_H

# Azure Core Skills
add_rect(slide, 0, y, LC, 0.55, AZURE)
add_textbox(slide, 0.25, y + 0.05, LC - 0.3, 0.5, 'Azure Core Skills', 9.5, bold=True, color=WHITE)
y += 0.55
add_bullet_box(slide, 0.25, y, LC - 0.4, 3.6, [
    'Azure Landing Zone (Open Cloud Framework)',
    'Hybrid Cloud Design & Implementation',
    'Azure Migration (DC → Cloud, eCommerce)',
    'Azure Optimisation & Cost Redesign',
    'Azure Virtual Network & Networking',
    'Azure IAM, RBAC & Security Centre',
    'Azure Monitor & Log Analytics',
    'Azure Backup & Disaster Recovery',
    'Azure Virtual Desktop / VDI',
    'Azure Policy & Governance',
], size=8.5)
y += 3.7

# DevOps & Automation
add_rect(slide, 0, y, LC, 0.55, AZURE)
add_textbox(slide, 0.25, y + 0.05, LC - 0.3, 0.5, 'DevOps & Automation', 9.5, bold=True, color=WHITE)
y += 0.55
add_bullet_box(slide, 0.25, y, LC - 0.4, 2.4, [
    'Terraform IaC (Azure & multi-cloud)',
    'Azure DevOps Pipelines',
    'Buildkite CI/CD pipelines',
    'GitHub Actions workflows',
    'SCCM / OS deployment automation',
], size=8.5)
y += 2.5

# Certifications
add_rect(slide, 0, y, LC, 0.55, NAVY)
add_textbox(slide, 0.25, y + 0.05, LC - 0.3, 0.5, 'Azure Certifications', 9.5, bold=True, color=WHITE)
y += 0.55
add_bullet_box(slide, 0.25, y, LC - 0.4, H - y - 0.1, [
    'MS: Azure Solutions Architect Expert',
    'MS: DevOps Engineer Expert',
    'MS: Azure Solutions Architect Design',
    'MCSE: Cloud Platform & Infrastructure',
    'MCSA: Cloud Platform',
    'MS: Server Virtualisation (Hyper-V)',
    'AWS: Solutions Architect Associate',
    'Google Cloud: Assoc. Cloud Engineer',
    'ITIL V4  ·  ScrumMaster®',
], size=8.2)

# ── RIGHT COLUMN ──────────────────────────────────────────────────────────────
y_r = HEADER_H
add_rect(slide, RC_X, y_r, RC_W, 0.55, AZURE)
add_textbox(slide, RC_X + 0.2, y_r + 0.05, RC_W - 0.3, 0.5,
            'Azure Project Portfolio', 9.5, bold=True, color=WHITE)
y_r += 0.55

projects = [
    (
        'Global Insurance Broker — Azure Migration  |  Accenture',
        'Data Centre Lead & Azure SME',
        '£1.2M savings  ·  UK-first financial cloud migration  ·  E2E delivery lead',
        [
            'Led UK-first large-scale migration of financial apps & infrastructure to Azure',
            'Delivery lead for ecommerce DC-to-Azure migration; Finance Transformation architect',
            'Azure Optimisation & Redesign — reduced cloud spend by ~£1.2 million',
        ]
    ),
    (
        'Global Manufacturing — Azure Landing Zone & Hybrid Cloud  |  Deloitte',
        'Azure Workstream Lead & Cloud Architect',
        'Enterprise Landing Zone  ·  Global multi-site  ·  Terraform + Buildkite',
        [
            'Architected enterprise-scale Azure landing zone (Open Cloud Framework) — global BU sites',
            'Terraform IaC deployed via Buildkite pipelines for repeatable, auditable delivery',
        ]
    ),
    (
        'UAE Retail Bank — Digital Banking Platform  |  Deloitte',
        'Azure Infrastructure DevOps Engineer',
        'Financial Services  ·  Regulated environment  ·  Terraform + Azure DevOps',
        [
            'Designed & implemented Azure cloud infrastructure for new digital banking platform',
        ]
    ),
    (
        'Market Platform Provider — Digital Market Platform  |  Deloitte',
        'Azure Infrastructure DevOps Engineer',
        'Azure DevOps  ·  Platform engineering  ·  Pipeline automation',
        [
            'Azure Infrastructure DevOps for Digital Market Platform Support & automation',
        ]
    ),
    (
        'UK Communications & Media — Cisco WebEx Rollout  |  Accenture',
        'Technology Lead',
        '15,000-user enterprise rollout  ·  Cross-functional delivery',
        [
            'Led technical release & enterprise-wide adoption of Cisco WebEx to 15,000 users',
        ]
    ),
    (
        'UK Retail Bank — Cloud Enablement & Architecture  |  Accenture',
        'Delivery Lead & Cloud Architect',
        'Cloud-native reference architectures  ·  Stakeholder enablement',
        [
            'Developed cloud-native reference architectures; worked with client architects on Azure strategy',
        ]
    ),
]

for title, role, badges, bullets in projects:
    if y_r > H - 1.0:
        break
    add_textbox(slide, RC_X + 0.15, y_r, RC_W - 0.2, 0.45, title, 9, bold=True, color=NAVY)
    y_r += 0.43
    add_textbox(slide, RC_X + 0.15, y_r, RC_W - 0.2, 0.32, role, 8.2, color=AZURE, italic=True)
    y_r += 0.30
    add_textbox(slide, RC_X + 0.15, y_r, RC_W - 0.2, 0.32, badges, 7.8, color=GREEN, italic=True)
    y_r += 0.30
    for b in bullets:
        add_textbox(slide, RC_X + 0.3, y_r, RC_W - 0.4, 0.34, f'▸  {b}', 8.2, color=RGBColor(0x1a, 0x1a, 0x1a))
        y_r += 0.32
    y_r += 0.18

# ── FOOTER ────────────────────────────────────────────────────────────────────
footer_y = H - 0.55
add_rect(slide, 0, footer_y, W, 0.55, NAVY)
add_textbox(slide, 0.4, footer_y + 0.07, 15, 0.45,
            'Deloitte  2020–Present  ·  Accenture  2016–2020  ·  Avanade  2011–2015  ·  Barclays Capital  2007–2011',
            7.5, color=WHITE)
add_textbox(slide, 16, footer_y + 0.07, 9.2, 0.45,
            'MCA First Class  ·  BSc Maths First Class  ·  School Governor 2018–present',
            7.5, color=LTBLUE, align=PP_ALIGN.RIGHT)

output = '/Users/karthiksankaran/PS-Scripts/CV/Karthik_Sankaran_OnePager_Azure_2026.pptx'
prs.save(output)
print(f'Saved: {output}')