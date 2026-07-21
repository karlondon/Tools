import openpyxl
import re
import json
from collections import Counter, defaultdict
from datetime import datetime

wb = openpyxl.load_workbook('task.xlsx')
ws = wb['Page 1']

rows = []
for r in range(2, ws.max_row + 1):
    rows.append({
        'Number': ws.cell(r, 1).value,
        'Priority': str(ws.cell(r, 2).value or ''),
        'State': str(ws.cell(r, 3).value or ''),
        'Assignment group': str(ws.cell(r, 4).value or ''),
        'Assigned to': str(ws.cell(r, 5).value or ''),
        'Short description': str(ws.cell(r, 6).value or ''),
        'Task type': str(ws.cell(r, 7).value or ''),
        'Updated': ws.cell(r, 8).value,
        'Created': ws.cell(r, 9).value,
    })

pcode_pat = re.compile(r'[Pp]\d{3,6}')
TARGET_GROUPS = {'GCP Engineers', 'Service Ops Engineers', 'Solutions Operations'}

for r in rows:
    r['pcodes'] = sorted(set(m.upper() for m in pcode_pat.findall(r['Short description'])))

pcode_tickets = [r for r in rows if r['pcodes']]
team_tickets = [r for r in rows if r['Assignment group'] in TARGET_GROUPS]

all_group_counts = Counter(r['Assignment group'] for r in rows if r['Assignment group'])
all_pcodes = sorted(set(p for r in pcode_tickets for p in r['pcodes']))
tickets_per_pcode = Counter(p for r in pcode_tickets for p in r['pcodes'])
team_tickets_per_pcode = Counter(p for r in pcode_tickets if r['Assignment group'] in TARGET_GROUPS for p in r['pcodes'])

pcode_group_breakdown = defaultdict(Counter)
for r in pcode_tickets:
    for p in r['pcodes']:
        pcode_group_breakdown[p][r['Assignment group']] += 1

task_type_counts = Counter(r['Task type'] for r in rows if r['Task type'])
state_counts = Counter(r['State'] for r in rows if r['State'])

assignee_counts = Counter(r['Assigned to'] for r in rows if r['Assigned to'])

monthly = Counter()
team_monthly = Counter()
pcode_monthly = Counter()
for r in rows:
    if r['Created']:
        key = r['Created'].strftime('%Y-%m')
        monthly[key] += 1
        if r['Assignment group'] in TARGET_GROUPS:
            team_monthly[key] += 1
        if r['pcodes']:
            pcode_monthly[key] += 1

sorted_months = sorted(monthly.keys())

# ---- HTML helpers ----
def bar(value, max_val, width=180, color='#1a73e8'):
    pct = int((value / max_val) * width) if max_val else 0
    return f'<span class="bar" style="width:{pct}px;background:{color}"></span>'

# ---- pCode table ----
pcode_rows_html = ''
max_p = max(tickets_per_pcode.values()) if tickets_per_pcode else 1
for p in sorted(tickets_per_pcode, key=lambda x: -tickets_per_pcode[x]):
    total = tickets_per_pcode[p]
    team_c = team_tickets_per_pcode.get(p, 0)
    groups = ', '.join(f'{g} ({c})' for g, c in sorted(pcode_group_breakdown[p].items(), key=lambda x: -x[1]))
    pcode_rows_html += f"""<tr>
      <td><span class="badge">{p}</span></td>
      <td>{total} {bar(total, max_p)}</td>
      <td>{team_c}</td>
      <td style="font-size:.82rem;color:#555">{groups}</td>
    </tr>"""

# ---- group table ----
group_rows_html = ''
max_g = max(all_group_counts.values()) if all_group_counts else 1
for g, c in sorted(all_group_counts.items(), key=lambda x: -x[1]):
    hl = ' style="color:#1a73e8;font-weight:600"' if g in TARGET_GROUPS else ''
    group_rows_html += f'<tr><td{hl}>{g}</td><td>{c} {bar(c, max_g)}</td></tr>'

# ---- assignee table ----
assignee_rows_html = ''
max_a = max(assignee_counts.values()) if assignee_counts else 1
for a, c in sorted(assignee_counts.items(), key=lambda x: -x[1])[:15]:
    assignee_rows_html += f'<tr><td>{a}</td><td>{c} {bar(c, max_a, 120, "#34a853")}</td></tr>'

# ---- state table ----
state_rows_html = ''
for s, c in sorted(state_counts.items(), key=lambda x: -x[1]):
    state_rows_html += f'<tr><td>{s}</td><td>{c}</td></tr>'

# ---- pCode ticket detail table ----
detail_rows_html = ''
for r in sorted(pcode_tickets, key=lambda x: str(x['Number']), reverse=True):
    badges = ' '.join(f'<span class="badge">{p}</span>' for p in r['pcodes'])
    hl = 'team-row' if r['Assignment group'] in TARGET_GROUPS else ''
    created = r['Created'].strftime('%Y-%m-%d') if r['Created'] else ''
    detail_rows_html += f"""<tr class="{hl}">
      <td>{r['Number']}</td>
      <td>{badges}</td>
      <td>{r['Assignment group']}</td>
      <td>{r['Assigned to']}</td>
      <td>{r['State']}</td>
      <td>{r['Task type']}</td>
      <td style="max-width:300px">{r['Short description'][:100]}</td>
      <td>{created}</td>
    </tr>"""

# ---- All tickets table ----
all_rows_html = ''
for r in sorted(rows, key=lambda x: str(x['Number']), reverse=True):
    badges = ' '.join(f'<span class="badge">{p}</span>' for p in r['pcodes']) if r['pcodes'] else ''
    hl = 'team-row' if r['Assignment group'] in TARGET_GROUPS else ''
    created = r['Created'].strftime('%Y-%m-%d') if r['Created'] else ''
    all_rows_html += f"""<tr class="{hl}">
      <td>{r['Number']}</td>
      <td>{badges if badges else '<span style="color:#aaa">—</span>'}</td>
      <td>{r['Assignment group']}</td>
      <td>{r['Assigned to']}</td>
      <td>{r['State']}</td>
      <td>{r['Task type']}</td>
      <td style="max-width:300px">{r['Short description'][:100]}</td>
      <td>{created}</td>
    </tr>"""

gen_time = datetime.now().strftime('%d %B %Y %H:%M')
unique_pcode_count = len(all_pcodes)
pcode_list_badges = ' '.join(f'<span class="badge badge-green">{p}</span>' for p in all_pcodes)

task_labels = list(task_type_counts.keys())
task_vals = list(task_type_counts.values())
state_labels = list(state_counts.keys())
state_vals = list(state_counts.values())
month_all = [monthly.get(m, 0) for m in sorted_months]
month_team = [team_monthly.get(m, 0) for m in sorted_months]
month_pcode_data = [pcode_monthly.get(m, 0) for m in sorted_months]

html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>GCP Engineers — Task Ticket Analysis</title>
<style>
*{{box-sizing:border-box;margin:0;padding:0}}
body{{font-family:'Segoe UI',Arial,sans-serif;background:#f0f4fb;color:#222}}
header{{background:linear-gradient(135deg,#1a73e8 0%,#0d47a1 100%);color:#fff;padding:36px 48px}}
header h1{{font-size:1.9rem;font-weight:700;margin-bottom:6px}}
header p{{opacity:.82;font-size:.95rem}}
.container{{max-width:1280px;margin:0 auto;padding:32px 24px}}
.kpi-grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:14px;margin-bottom:32px}}
.kpi{{background:#fff;border-radius:12px;padding:22px 18px;box-shadow:0 2px 10px rgba(0,0,0,.07);text-align:center}}
.kpi .num{{font-size:2.6rem;font-weight:800;color:#1a73e8;line-height:1}}
.kpi .lbl{{font-size:.78rem;color:#777;margin-top:6px;text-transform:uppercase;letter-spacing:.5px}}
.section{{background:#fff;border-radius:12px;padding:26px 28px;margin-bottom:26px;box-shadow:0 2px 10px rgba(0,0,0,.07)}}
.section h2{{font-size:1rem;color:#1a73e8;border-bottom:2px solid #e8f0fe;padding-bottom:10px;margin-bottom:18px;display:flex;align-items:center;gap:8px}}
table{{width:100%;border-collapse:collapse;font-size:.86rem}}
th{{background:#e8f0fe;color:#1565c0;padding:9px 11px;text-align:left;font-weight:600}}
td{{padding:7px 11px;border-bottom:1px solid #f0f0f0;vertical-align:middle}}
tr:hover td{{background:#f5f8ff}}
.team-row td{{background:#edf7ed}}
.team-row:hover td{{background:#d6f0d6}}
.badge{{display:inline-block;background:#1a73e8;color:#fff;border-radius:20px;padding:2px 10px;font-size:.75rem;margin:1px;font-weight:500}}
.badge-green{{background:#34a853}}
.bar{{display:inline-block;height:13px;border-radius:3px;vertical-align:middle;margin-left:6px}}
.chart-wrap{{position:relative;height:260px}}
.two-col{{display:grid;grid-template-columns:1fr 1fr;gap:22px}}
.pcode-list{{margin-bottom:14px;line-height:2}}
.note{{font-size:.78rem;color:#999;margin-top:8px}}
@media(max-width:700px){{.two-col{{grid-template-columns:1fr}}header{{padding:24px 20px}}.container{{padding:16px 12px}}}}
</style>
<script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.0/dist/chart.umd.min.js"></script>
</head>
<body>
<header>
  <h1>📊 GCP Engineers — Task Ticket Analysis</h1>
  <p>Source: task.xlsx &nbsp;|&nbsp; Generated {gen_time} &nbsp;|&nbsp; Covering {sorted_months[0] if sorted_months else 'N/A'} → {sorted_months[-1] if sorted_months else 'N/A'}</p>
</header>
<div class="container">

<!-- KPI Cards -->
<div class="kpi-grid">
  <div class="kpi"><div class="num">{len(rows)}</div><div class="lbl">Total Tickets</div></div>
  <div class="kpi"><div class="num">{len(pcode_tickets)}</div><div class="lbl">Tickets with pCode</div></div>
  <div class="kpi"><div class="num">{unique_pcode_count}</div><div class="lbl">Unique pCodes (Clients)</div></div>
  <div class="kpi"><div class="num">{len(team_tickets)}</div><div class="lbl">GCP Eng / SolOps</div></div>
  <div class="kpi"><div class="num">{len(sorted_months)}</div><div class="lbl">Months Active</div></div>
  <div class="kpi"><div class="num">{len(assignee_counts)}</div><div class="lbl">Engineers Involved</div></div>
</div>

<!-- pCode Summary -->
<div class="section">
  <h2>🏷️ Unique Client pCodes Identified — {unique_pcode_count} Clients</h2>
  <div class="pcode-list">{pcode_list_badges}</div>
  <p class="note">pCodes extracted from ticket short descriptions (pattern: P + 3–6 digits). Each pCode represents a distinct client/project.</p>
</div>

<!-- Monthly Trend -->
<div class="section">
  <h2>📈 Monthly Ticket Volume Trend</h2>
  <div class="chart-wrap"><canvas id="monthChart"></canvas></div>
  <p class="note">Blue bars = all tickets · Green line = GCP Eng / SolOps team · Red line = tickets referencing a pCode</p>
</div>

<!-- pCode Analysis -->
<div class="section">
  <h2>🔢 Ticket Count per pCode (Client)</h2>
  <table>
    <thead><tr><th>pCode</th><th>Total Tickets</th><th>Team Tickets</th><th>Handled By</th></tr></thead>
    <tbody>{pcode_rows_html}</tbody>
  </table>
</div>

<!-- Charts row -->
<div class="two-col">
  <div class="section">
    <h2>📋 Task Type Distribution</h2>
    <div class="chart-wrap" style="height:220px"><canvas id="taskChart"></canvas></div>
  </div>
  <div class="section">
    <h2>✅ Ticket State Breakdown</h2>
    <div class="chart-wrap" style="height:220px"><canvas id="stateChart"></canvas></div>
  </div>
</div>

<!-- Assignment Group -->
<div class="section">
  <h2>👥 Tickets by Assignment Group</h2>
  <table>
    <thead><tr><th>Assignment Group</th><th>Ticket Count</th></tr></thead>
    <tbody>{group_rows_html}</tbody>
  </table>
  <p class="note">Blue = target teams (GCP Engineers, Service Ops Engineers, Solutions Operations)</p>
</div>

<!-- Top Assignees -->
<div class="section">
  <h2>🧑‍💻 Top Engineers by Ticket Count</h2>
  <table>
    <thead><tr><th>Engineer</th><th>Tickets</th></tr></thead>
    <tbody>{assignee_rows_html}</tbody>
  </table>
</div>

<!-- pCode tickets detail -->
<div class="section">
  <h2>🔍 pCode Tickets Detail ({len(pcode_tickets)} tickets — green = target team)</h2>
  <table>
    <thead><tr><th>Ticket #</th><th>pCode(s)</th><th>Group</th><th>Assigned To</th><th>State</th><th>Type</th><th>Description</th><th>Created</th></tr></thead>
    <tbody>{detail_rows_html}</tbody>
  </table>
</div>

<!-- All tickets -->
<div class="section">
  <h2>📄 All Tickets ({len(rows)} total)</h2>
  <table>
    <thead><tr><th>Ticket #</th><th>pCode(s)</th><th>Group</th><th>Assigned To</th><th>State</th><th>Type</th><th>Description</th><th>Created</th></tr></thead>
    <tbody>{all_rows_html}</tbody>
  </table>
</div>

</div>
<script>
const months = {json.dumps(sorted_months)};
const allData = {json.dumps(month_all)};
const teamData = {json.dumps(month_team)};
const pcodeData = {json.dumps(month_pcode_data)};

new Chart(document.getElementById('monthChart'), {{
  type: 'bar',
  data: {{
    labels: months,
    datasets: [
      {{label:'All Tickets',data:allData,backgroundColor:'rgba(26,115,232,0.22)',borderColor:'#1a73e8',borderWidth:1.5,order:2}},
      {{label:'GCP Eng / SolOps',data:teamData,type:'line',borderColor:'#34a853',backgroundColor:'rgba(52,168,83,0.1)',borderWidth:2,pointRadius:5,fill:true,tension:0.3,order:1}},
      {{label:'pCode Tickets',data:pcodeData,type:'line',borderColor:'#ea4335',backgroundColor:'transparent',borderWidth:2,pointRadius:5,fill:false,tension:0.3,order:0}}
    ]
  }},
  options:{{responsive:true,maintainAspectRatio:false,plugins:{{legend:{{position:'top'}}}},scales:{{y:{{beginAtZero:true,ticks:{{stepSize:1}}}}}}}}
}});

new Chart(document.getElementById('taskChart'), {{
  type: 'doughnut',
  data: {{
    labels: {json.dumps(task_labels)},
    datasets: [{{data:{json.dumps(task_vals)},backgroundColor:['#1a73e8','#34a853','#fbbc04','#ea4335','#9c27b0','#00bcd4','#ff7043']}}]
  }},
  options:{{responsive:true,maintainAspectRatio:false,plugins:{{legend:{{position:'bottom',labels:{{font:{{size:11}}}}}}}}}}
}});

new Chart(document.getElementById('stateChart'), {{
  type: 'doughnut',
  data: {{
    labels: {json.dumps(state_labels)},
    datasets: [{{data:{json.dumps(state_vals)},backgroundColor:['#34a853','#1a73e8','#fbbc04','#ea4335','#9e9e9e','#00bcd4']}}]
  }},
  options:{{responsive:true,maintainAspectRatio:false,plugins:{{legend:{{position:'bottom',labels:{{font:{{size:11}}}}}}}}}}
}});
</script>
</body>
</html>"""

with open('GCP_Tasks_Analysis.html', 'w') as f:
    f.write(html)

print("Report written to GCP_Tasks_Analysis.html")
print(f"Total: {len(rows)} | pCode tickets: {len(pcode_tickets)} | Unique pCodes: {unique_pcode_count}")