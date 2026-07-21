import openpyxl
import re
import json
from collections import Counter, defaultdict
from datetime import datetime

wb = openpyxl.load_workbook('task.xlsx')
ws = wb['Page 1']

rows = []
for r in range(2, ws.max_row + 1):
    rows.append({
        'Number': ws.cell(r, 1).value,
        'Priority': str(ws.cell(r, 2).value or ''),
        'State': str(ws.cell(r, 3).value or ''),
        'Assignment group': str(ws.cell(r, 4).value or ''),
        'Assigned to': str(ws.cell(r, 5).value or ''),
        'Short description': str(ws.cell(r, 6).value or ''),
        'Task type': str(ws.cell(r, 7).value or ''),
        'Updated': ws.cell(r, 8).value,
        'Created': ws.cell(r, 9).value,
    })

pcode_pat = re.compile(r'[Pp]\d{3,6}')
TARGET_GROUPS = {'GCP Engineers', 'Service Ops Engineers', 'Solutions Operations'}

for r in rows:
    pcodes = sorted(set(m.upper() for m in pcode_pat.findall(r['Short description'])))
    r['pcodes'] = pcodes

pcode_tickets = [r for r in rows if r['pcodes']]
team_tickets = [r for r in rows if r['Assignment group'] in TARGET_GROUPS]
team_pcode_tickets = [r for r in team_tickets if r['pcodes']]

all_group_counts = Counter(r['Assignment group'] for r in rows if r['Assignment group'])
all_pcodes = sorted(set(p for r in pcode_tickets for p in r['pcodes']))
tickets_per_pcode = Counter(p for r in pcode_tickets for p in r['pcodes'])
team_tickets_per_pcode = Counter(p for r in team_pcode_tickets for p in r['pcodes'])

pcode_group_breakdown = defaultdict(Counter)
for r in pcode_tickets:
    for p in r['pcodes']:
        pcode_group_breakdown[p][r['Assignment group']] += 1

task_type_counts = Counter(r['Task type'] for r in rows if r['Task type'])
state_counts = Counter(r['State'] for r in rows if r['State'])

monthly = Counter()
team_monthly = Counter()
pcode_monthly = Counter()
for r in rows:
    if r['Created']:
        key = r['Created'].strftime('%Y-%m')
        monthly[key] += 1
        if r['Assignment group'] in TARGET_GROUPS:
            team_monthly[key] += 1
        if r['pcodes']:
            pcode_monthly[key] += 1

sorted_months = sorted(monthly.keys())

# Save stats to JSON for the report
stats = {
    'total': len(rows),
    'pcode_tickets': len(pcode_tickets),
    'unique_pcodes': len(all_pcodes),
    'all_pcodes': all_pcodes,
    'team_total': sum(1 for r in rows if r['Assignment group'] in TARGET_GROUPS),
    'team_pcode_tickets': len(team_pcode_tickets),
    'months': len(sorted_months),
    'group_counts': dict(all_group_counts),
    'tickets_per_pcode': dict(tickets_per_pcode),
    'team_tickets_per_pcode': dict(team_tickets_per_pcode),
    'pcode_group_breakdown': {p: dict(v) for p, v in pcode_group_breakdown.items()},
    'task_type_counts': dict(task_type_counts),
    'state_counts': dict(state_counts),
    'sorted_months': sorted_months,
    'month_all': [monthly.get(m, 0) for m in sorted_months],
    'month_team': [team_monthly.get(m, 0) for m in sorted_months],
    'month_pcode': [pcode_monthly.get(m, 0) for m in sorted_months],
}

with open('/tmp/gcp_tasks_stats.json', 'w') as f:
    json.dump(stats, f)

# Print summary
print(f"Total tickets: {stats['total']}")
print(f"Tickets with pCode: {stats['pcode_tickets']}")
print(f"Unique pCodes (clients): {stats['unique_pcodes']}")
print(f"All pCodes: {all_pcodes}")
print(f"Team tickets: {stats['team_total']}")
print(f"Assignment groups: {dict(all_group_counts)}")
print(f"Months: {sorted_months}")
print(f"pCode ticket sample:")
for r in pcode_tickets[:10]:
    print(f"  {r['Number']} | {r['pcodes']} | {r['Assignment group']} | {r['Short description'][:70]}")
print("DONE")